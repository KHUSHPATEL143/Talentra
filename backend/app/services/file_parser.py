"""File parsing helpers for PDF, DOCX, and text resumes."""

from __future__ import annotations

from io import BytesIO

import fitz
import pdfplumber
from docx import Document


class FileParserService:
    """Extract raw text from supported resume document formats."""

    async def extract_text(self, file_bytes: bytes, mime_type: str) -> str:
        """Dispatch to the appropriate parser based on MIME type."""

        if mime_type == "application/pdf":
            return self._extract_pdf_text(file_bytes)
        if mime_type in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        }:
            return self._extract_docx_text(file_bytes)
        if mime_type.startswith("text/"):
            return self._extract_text_file(file_bytes)
        raise ValueError(f"Unsupported MIME type: {mime_type}")

    def _extract_pdf_text(self, file_bytes: bytes) -> str:
        """Extract PDF text with PyMuPDF and fallback to pdfplumber for sparse layouts."""

        document = fitz.open(stream=file_bytes, filetype="pdf")
        pages = []
        for page in document:
            blocks = page.get_text("blocks")
            page_text = "\n".join(block[4].strip() for block in blocks if len(block) > 4 and block[4].strip())
            if page_text:
                pages.append(page_text)
        extracted = "\n\n".join(pages).strip()
        if extracted:
            return extracted

        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            plumber_pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n\n".join(page.strip() for page in plumber_pages if page.strip())

    def _extract_docx_text(self, file_bytes: bytes) -> str:
        """Extract DOCX paragraph and table content."""

        document = Document(BytesIO(file_bytes))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts).strip()

    def _extract_text_file(self, file_bytes: bytes) -> str:
        """Decode text content using UTF-8 with replacement."""

        return file_bytes.decode("utf-8", errors="replace").strip()
